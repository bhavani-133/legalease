import io
import requests
import streamlit as st

from docx import Document
from fpdf import FPDF
# ---------------------------------------------------------
# DOCUMENT EXPORT FUNCTIONS
# ---------------------------------------------------------

def create_docx(document_text):
    """Create a real DOCX file."""

    doc = Document()

    # Add LegalEase title
    title = doc.add_paragraph()
    title_run = title.add_run("LegalEase")
    title_run.bold = True

    # Add document content
    for paragraph in document_text.split("\n\n"):
        paragraph = paragraph.strip()

        if paragraph:
            doc.add_paragraph(paragraph)

    # Save DOCX into memory
    file_stream = io.BytesIO()
    doc.save(file_stream)

    file_stream.seek(0)

    return file_stream.getvalue()


def create_pdf(document_text):
    """Create a real PDF file."""

    pdf = FPDF()

    pdf.set_auto_page_break(
        auto=True,
        margin=20
    )

    pdf.add_page()

    # LegalEase title
    pdf.set_font("Helvetica", "B", 16)

    pdf.cell(
        0,
        10,
        "LegalEase",
        new_x="LMARGIN",
        new_y="NEXT",
        align="C"
    )

    pdf.ln(5)

    # Document text
    pdf.set_font("Helvetica", size=11)

    # Handle special characters
    safe_text = (
        document_text
        .replace("–", "-")
        .replace("—", "-")
        .replace("“", '"')
        .replace("”", '"')
        .replace("‘", "'")
        .replace("’", "'")
        .replace("•", "-")
    )

    for paragraph in safe_text.split("\n\n"):

        paragraph = paragraph.strip()

        if paragraph:
            pdf.multi_cell(
                0,
                7,
                paragraph
            )

            pdf.ln(3)

    # Save PDF into memory
    pdf_stream = io.BytesIO()

    pdf.output(pdf_stream)

    return pdf_stream.getvalue()


# ---------------------------------------------------------
# PAGE CONFIGURATION
# ---------------------------------------------------------

st.set_page_config(
    page_title="LegalEase - AI Legal Document Generator",
    page_icon="⚖️",
    layout="wide"
)


# ---------------------------------------------------------
# CUSTOM CSS
# ---------------------------------------------------------

st.markdown(
    """
    <style>

    .main-title {
        text-align: center;
        font-size: 42px;
        font-weight: 700;
        margin-bottom: 5px;
    }

    .subtitle {
        text-align: center;
        font-size: 18px;
        margin-bottom: 30px;
    }

    .section-title {
        font-size: 22px;
        font-weight: 600;
        margin-top: 20px;
        margin-bottom: 10px;
    }

    .legal-note {
        padding: 15px;
        border-radius: 8px;
        border: 1px solid #ddd;
        margin-top: 20px;
        font-size: 14px;
    }

    </style>
    """,
    unsafe_allow_html=True
)


# ---------------------------------------------------------
# TITLE
# ---------------------------------------------------------

st.markdown(
    '<div class="main-title">⚖️ LegalEase</div>',
    unsafe_allow_html=True
)

st.markdown(
    '<div class="subtitle">AI-Powered Legal Document Generator</div>',
    unsafe_allow_html=True
)


st.divider()


# ---------------------------------------------------------
# BACKEND URL
# ---------------------------------------------------------

BACKEND_URL = "http://127.0.0.1:8000"


# ---------------------------------------------------------
# INPUT SECTION
# ---------------------------------------------------------

st.markdown(
    '<div class="section-title">1. Document Type</div>',
    unsafe_allow_html=True
)

document_type = st.selectbox(
    "Select the type of legal document",
    [
        "Freelance Work Contract",
        "Employment Contract",
        "Non-Disclosure Agreement (NDA)",
        "Lease Agreement",
        "Service Agreement",
        "Other"
    ]
)


st.markdown(
    '<div class="section-title">2. Parties Involved</div>',
    unsafe_allow_html=True
)

parties = st.text_area(
    "Enter the names and roles of the parties",
    placeholder=(
        "Example: Jane Doe (Service Provider), "
        "TechNova Inc. (Client)"
    ),
    height=100
)


st.markdown(
    '<div class="section-title">3. Terms & Conditions</div>',
    unsafe_allow_html=True
)

terms = st.text_area(
    "Enter the important terms and conditions",
    placeholder=(
        "Separate each term using a semicolon (;)\n\n"
        "Example:\n"
        "Payment within 30 days of invoice; "
        "Work must be completed by the agreed deadline; "
        "Confidentiality must be maintained; "
        "Either party may terminate with 15 days notice"
    ),
    height=160
)


st.markdown(
    '<div class="section-title">4. Effective Date</div>',
    unsafe_allow_html=True
)

effective_date = st.date_input(
    "Select the effective date"
)


st.write("")


# ---------------------------------------------------------
# GENERATE BUTTON
# ---------------------------------------------------------

generate_button = st.button(
    "Generate Document",
    type="primary",
    use_container_width=True
)


# ---------------------------------------------------------
# DOCUMENT GENERATION
# ---------------------------------------------------------

if generate_button:

    if not parties.strip():
        st.error("Please enter the parties involved.")

    elif not terms.strip():
        st.error("Please enter the terms and conditions.")

    else:

        request_data = {
            "document_type": document_type,
            "parties": parties,
            "terms": terms,
            "effective_date": effective_date.strftime("%B %d, %Y")
        }

        try:

            with st.spinner("Generating your legal document..."):

                response = requests.post(
                    f"{BACKEND_URL}/generate",
                    json=request_data,
                    timeout=60
                )

            if response.status_code == 200:

                result = response.json()

                if result.get("success"):

                    st.session_state["document"] = result["document"]

                    st.success(
                        "Your legal document has been generated successfully!"
                    )

                else:

                    st.error(
                        "The document could not be generated."
                    )

            else:

                try:
                    error_message = response.json().get(
                        "detail",
                        "Unknown backend error."
                    )
                except Exception:
                    error_message = response.text

                st.error(
                    f"Backend error ({response.status_code}): "
                    f"{error_message}"
                )

        except requests.exceptions.ConnectionError:

            st.error(
                "Could not connect to the FastAPI backend. "
                "Make sure the backend is running on "
                "http://127.0.0.1:8000"
            )

        except requests.exceptions.Timeout:

            st.error(
                "The request took too long. "
                "Please try again."
            )

        except Exception as error:

            st.error(
                f"Unexpected error: {error}"
            )


# ---------------------------------------------------------
# GENERATED DOCUMENT
# ---------------------------------------------------------

if "document" in st.session_state:

    st.divider()

    st.markdown(
        '<div class="section-title">Generated Document</div>',
        unsafe_allow_html=True
    )

    edited_document = st.text_area(
        "Click inside the document below to edit it:",
        value=st.session_state["document"],
        height=600
    )

    # Save any edits
    st.session_state["document"] = edited_document


    # -----------------------------------------------------
# DOWNLOAD BUTTONS
# -----------------------------------------------------

st.markdown(
    '<div class="section-title">Download / Save</div>',
    unsafe_allow_html=True
)

# Create real DOCX and PDF files
edited_document = st.session_state.get("document", "")
docx_file = create_docx(edited_document)
pdf_file = create_pdf(edited_document)

col1, col2, col3 = st.columns(3)


with col1:

    st.download_button(
        label="Download TXT",
        data=edited_document,
        file_name="LegalEase_Document.txt",
        mime="text/plain",
        use_container_width=True
    )


with col2:

    st.download_button(
        label="Download DOCX",
        data=docx_file,
        file_name="LegalEase_Document.docx",
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        use_container_width=True
    )


with col3:

    st.download_button(
        label="Download PDF",
        data=pdf_file,
        file_name="LegalEase_Document.pdf",
        mime="application/pdf",
        use_container_width=True
    )